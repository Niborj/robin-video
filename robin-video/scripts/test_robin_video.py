"""Robin-Video self-tests. Run:  uv run --project <skill-dir>/scripts python <skill-dir>/scripts/test_robin_video.py
Covers ingestion (PDF, PPTX with a chart, DOCX with a layout table, Markdown, HTML, CSV), the storyboard build in every
format and archetype, the brand gate (positive and negative), and the timing math. No rendering (that needs Chrome)."""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(HERE))
import brand_check  # noqa: E402
import kit  # noqa: E402
import scenes  # noqa: E402
import voice  # noqa: E402

SAMPLE_PDF = SKILL / 'examples/sample-source/sample-ai-adoption-brief.pdf'
GALLERY = SKILL / 'examples/gallery/storyboard.json'


def ingest(files, out):
    r = subprocess.run([sys.executable, str(HERE / 'ingest.py'), *map(str, files), '--out', str(out)], capture_output=True, text=True)
    if r.returncode != 0:
        raise AssertionError(r.stderr or r.stdout)
    return json.loads((out / 'content.json').read_text())


class Ingest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_pdf(self):
        c = ingest([SAMPLE_PDF], self.tmp / 'pdf')
        self.assertIn('AI adoption', c['title'])
        values = [s['value'] for s in c['stats'] if s['strength'] == 'strong']
        self.assertIn('71%', values)
        self.assertIn('30,254', values)
        self.assertTrue(any('Visibility came first' in q['text'] for q in c['quotes']))
        self.assertEqual(len(c['pages']), 2)
        self.assertTrue((self.tmp / 'pdf' / c['pages'][0]['image']).is_file())
        self.assertTrue(any(b['type'] == 'bullet' and 'See every AI app' in b['text'] for b in c['blocks']))
        gaps = c['pages'][1]['gaps']                                  # blank bands for doc framing
        self.assertGreater(len(gaps), 5)
        self.assertEqual(gaps, sorted(gaps))
        self.assertTrue(all(0 <= a < b <= 1 for a, b in gaps))

    def test_blank_bands(self):
        from PIL import Image, ImageDraw
        im = Image.new('RGB', (200, 1000), 'white')
        d = ImageDraw.Draw(im)
        d.rectangle([20, 100, 180, 199], fill='black')               # a "line" of text from y 100 to 199
        d.rectangle([20, 400, 180, 409], fill='black')
        path = self.tmp / 'bands.png'; im.save(path)
        import ingest as I
        self.assertEqual(I.blank_bands(path), [[0.0, 0.1], [0.2, 0.4], [0.41, 1.0]])

    def test_pptx_chart_and_notes(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = 'Agentic attacks are rising'
        s.placeholders[1].text_frame.text = 'Attackers automate reconnaissance'
        s.notes_slide.notes_text_frame.text = 'Mention the 3x increase.'
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        s2.shapes.title.text = 'Share of traffic'
        data = CategoryChartData(); data.categories = ['Web AI', 'Agents']; data.add_series('2026', (62, 21))
        s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), data)
        path = self.tmp / 'deck.pptx'; prs.save(path)
        c = ingest([path], self.tmp / 'pptx')
        self.assertEqual(c['title'], 'Agentic attacks are rising')
        self.assertEqual(c['charts'][0]['categories'], ['Web AI', 'Agents'])
        self.assertEqual(c['charts'][0]['series'][0]['values'], [62.0, 21.0])
        self.assertTrue(any(b['type'] == 'notes' for b in c['blocks']))

    def test_docx_layout_table(self):
        import docx
        d = docx.Document()
        d.add_heading('Field notes on AI', 0)
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = 'Left column with a long paragraph that is clearly content rather than a data table, written in full sentences for the reader.'
        t.cell(0, 1).paragraphs[0].add_run('What changed').bold = True
        t.cell(0, 1).add_paragraph('• 94% of measured AI traffic needs governance review.')
        path = self.tmp / 'notes.docx'; d.save(path)
        c = ingest([path], self.tmp / 'docx')
        self.assertEqual(c['title'], 'Field notes on AI')
        types = {b['text'][:12]: b['type'] for b in c['blocks']}
        self.assertEqual(types.get('What changed'), 'heading')
        self.assertTrue(any(b['type'] == 'bullet' and b['text'].startswith('94%') for b in c['blocks']))
        self.assertFalse(c['tables'], 'layout tables are read as content, not data')

    def test_markdown_html_csv_and_flags(self):
        md = self.tmp / 'post.md'
        md.write_text('# Internal draft: AI rollout\n\nOnly 1 in 3 teams has an AI policy. Contact jane@example.com.\n\n- See\n- Decide\n')
        html = self.tmp / 'page.html'
        html.write_text('<html><title>T</title><body><h1>Visibility first</h1><p>We saw 3x more agent traffic this year than last year.</p><script>x</script></body></html>')
        csv = self.tmp / 'data.csv'
        csv.write_text('category,share\nweb,62\nagents,21\n')
        c = ingest([md, html, csv], self.tmp / 'mixed')
        self.assertEqual(c['title'], 'Internal draft: AI rollout')
        self.assertIn('draft', c['flags']['confidentialMarkers'])
        self.assertEqual(c['flags']['emails'][0]['value'], 'jane@example.com')
        self.assertTrue(any(s['value'].startswith('1 in 3') for s in c['stats']))
        self.assertTrue(any(s['value'] == '3x' for s in c['stats']))
        self.assertEqual(c['tables'][0]['rows'][1], ['web', '62'])
        self.assertIn('Clearance flags', (self.tmp / 'mixed/content.md').read_text())


class Build(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, sb, out=None):
        out = out or self.tmp / 'v'
        out.mkdir(parents=True, exist_ok=True)
        (out / 'source/pages').mkdir(parents=True, exist_ok=True)
        shutil.copy2(SKILL / 'examples/sample-source/page-sample.png', out / 'source/pages/p-02.png')
        (out / 'storyboard.json').write_text(json.dumps(sb))
        kit.main(['build', str(out)])
        return out

    def test_gallery_every_format(self):
        sb = json.loads(GALLERY.read_text())
        for fmt in kit.FORMATS:
            sb['format'] = fmt
            out = self.build(sb, self.tmp / fmt)
            comp = out / 'composition'
            index = (comp / 'index.html').read_text()
            self.assertIn(f'fmt-{fmt}', index)
            self.assertEqual(index.count('data-composition-src='), len(sb['scenes']))
            for s in sb['scenes']:
                html = (comp / 'compositions' / f'{s["id"]}.html').read_text()
                self.assertTrue(html.startswith('<template'), s['id'])
                self.assertIn(f'window.__timelines["{s["id"]}"]', html)
                self.assertIn(f'data-composition-id="{s["id"]}"', html)
                self.assertNotIn('TODO', html, s['id'])
            findings = brand_check.check(comp, None)
            errors = [f for f in findings if f['severity'] == 'error']
            self.assertEqual(errors, [], f'{fmt}: {errors}')

    def test_layout_overlaps_and_total(self):
        sb = {'format': 'portrait', 'overlap': 0.8, 'scenes': [
            {'id': 'a', 'archetype': 'hook', 'duration': 4}, {'id': 'b', 'archetype': 'statement', 'duration': 3, 'transition': 'cut'},
            {'id': 'c', 'archetype': 'end', 'duration': 3}]}
        rows, total = kit.layout(sb)
        self.assertEqual(total, 10)
        self.assertEqual([r['start'] for r in rows], [0, 4, 7])
        self.assertEqual([r['slot'] for r in rows], [4, 3.8, 3])   # a is cut into b (no overlap); b overlaps c's iris
        self.assertEqual(rows[0]['transition'], 'cut')

    def test_kept_scene_resyncs_timing(self):
        sb = {'format': 'square', 'scenes': [{'id': 'hook', 'archetype': 'hook', 'duration': 4, 'headline': 'One [idea]'},
                                             {'id': 'end', 'archetype': 'end', 'duration': 3, 'line': 'Go', 'cta': 'Read it'}]}
        out = self.build(sb)
        f = out / 'composition/compositions/hook.html'
        f.write_text(f.read_text().replace('One <span class="hl">idea</span>', 'Edited <span class="hl">by hand</span>'))
        sb['scenes'][0]['duration'] = 5
        self.build(sb, out)
        html = f.read_text()
        self.assertIn('Edited', html)
        self.assertIn('data-duration="5.85"', html)

    def test_storyboard_validation(self):
        with self.assertRaises(SystemExit):
            kit.validate({'scenes': [{'id': 'Bad Id', 'archetype': 'hook', 'duration': 3}]})
        with self.assertRaises(SystemExit):
            kit.validate({'scenes': [{'id': 'a', 'archetype': 'nope', 'duration': 3}]})
        with self.assertRaises(SystemExit):
            kit.validate({'scenes': [{'id': 'a', 'archetype': 'hook', 'duration': 0}]})

    def test_rich_and_values(self):
        self.assertEqual(scenes.rich('AI moves [fast] <b>'), 'AI moves <span class="hl">fast</span> &lt;b&gt;')
        self.assertEqual(scenes.split_value('$5.02M'), ('$', '5.02', 'M'))
        self.assertEqual(scenes.split_value('71%'), ('', '71', '%'))
        self.assertEqual(scenes.split_value('3x'), ('', '3', '×'))
        self.assertEqual(scenes.image_size(SKILL / 'examples/sample-source/page-sample.png')[0], 1600)


class DocFraming(unittest.TestCase):
    def ctx(self, **kw):
        return {'id': 'page', 'index': 0, 'fmt': 'portrait', 'w': 1080, 'h': 1350, 'pad_x': 84, 'bg': 'pine', 'img_size': (1600, 1131), **kw}

    def test_crop_to_content_and_zoom_on_focus(self):
        s = {'archetype': 'doc', 'image': 'assets/source/p.png', 'focus': [0.02, 0.23, 0.48, 0.195], 'tag': 'Look here'}
        whole = scenes.doc_frame(s, self.ctx())
        self.assertEqual((whole['page_w'], whole['window_h'], whole['y'], whole['h']), (912, 645, 0, 1))   # full page, no crop
        cropped = scenes.doc_frame(s, self.ctx(crop=[0, 0.0077, 1, 0.4576]))
        self.assertEqual((cropped['page_w'], cropped['window_h'], cropped['x']), (1629, 527, 0))     # zoomed around the focus (~2x)

    def test_window_keeps_proportions_instead_of_slicing(self):
        # The page is placed in percentages of a card with a fixed aspect ratio, so a shorter stage (burned-in
        # captions) scales the whole window down rather than cutting the page.
        s = {'archetype': 'doc', 'image': 'assets/source/p.png', 'focus': [0.02, 0.23, 0.48, 0.195]}
        html = scenes.a_doc(s, self.ctx())
        self.assertIn('height: min(100%, 645px); aspect-ratio: 912 / 645', html)
        self.assertIn('width: 100.00%; height: 100.00%; left: 0.00%; top: 0.00%', html)
        self.assertNotIn('clamp(', html)

    def test_tall_page_without_bands(self):
        s = {'archetype': 'doc', 'image': 'assets/source/p.png', 'focus': [0.1, 0.7, 0.5, 0.1]}
        ctx = self.ctx(img_size=(1600, 2263))                                  # up to 1/3 taller than the stage: whole, scaled
        d = scenes.doc_frame(s, ctx)
        self.assertEqual((d['y'], d['h']), (0, 1))
        self.assertGreaterEqual(scenes.stage_height(s, ctx) / d['window_h'], 0.75)
        ctx = self.ctx(img_size=(1600, 4000))                                  # much taller: a band centered on the focus
        d = scenes.doc_frame(s, ctx)
        self.assertLessEqual(d['y'], 0.7)
        self.assertGreaterEqual(d['y'] + d['h'], 0.8)
        self.assertAlmostEqual(d['window_h'], scenes.stage_height(s, ctx), delta=1)

    def square_page(self, **s):
        """The sample brief's p2 in square: a heading, a paragraph, a four-bar chart and a list, with blank bands."""
        gaps = [[0.0, 0.072], [0.102, 0.1236], [0.1373, 0.1444], [0.1581, 0.1656], [0.1792, 0.1868], [0.2004, 0.2724],
                [0.2887, 0.3276], [0.3563, 0.3987], [0.4274, 0.4698], [0.4989, 0.5413], [0.57, 0.649], [0.6658, 0.6958],
                [0.7095, 0.7289], [0.7426, 0.7625], [0.7762, 0.8221], [0.8499, 1.0]]
        ctx = {'id': 'page', 'index': 0, 'fmt': 'square', 'w': 1080, 'h': 1080, 'pad_x': 80, 'bg': 'pine', 'img_size': (1600, 2265),
               'crop': [0.0431, 0.047, 0.905, 0.8279], 'gaps': gaps}
        sc = {'archetype': 'doc', 'image': 'assets/source/p.png', 'headline': 'Where the AI traffic went', **s}
        return sc, ctx, gaps

    def test_window_edges_stay_in_blank_bands_through_the_push(self):
        sc, ctx, gaps = self.square_page(focus=[0.06, 0.27, 0.82, 0.25], tag='62% to web AI assistants')
        d = scenes.doc_frame(sc, ctx)
        top, bottom, zoom = d['y'], d['y'] + d['h'], 1.08
        o = 0.27 + 0.25 / 2
        blank = lambda y: any(a <= y <= b for a, b in gaps)
        for y in (top, bottom, o - (o - top) / zoom, o + (bottom - o) / zoom):   # at rest and at the end of the push
            self.assertTrue(blank(y), f'{y} cuts through a line ({d})')
        self.assertLessEqual(d['window_h'], scenes.stage_height(sc, ctx) + 1)
        self.assertLessEqual(top, 0.27 - 66 / d['page_h'])                     # room for the tag above the focus
        self.assertGreaterEqual(bottom, 0.52)
        self.assertEqual(d['tag'], 'above')                                   # its slot is blank space on the page

    def test_tag_goes_below_when_only_that_slot_is_blank(self):
        sc, ctx, gaps = self.square_page(focus=[0.06, 0.65, 0.5, 0.05], tag='See every AI app')   # list item; title right above
        ctx['gaps'] = [g for g in gaps if not g[0] < 0.62 < g[1]] + [[0.6, 0.625], [0.705, 0.75]]
        d = scenes.doc_frame(sc, ctx)
        self.assertEqual(d['tag'], 'below')
        self.assertGreaterEqual(d['y'] + d['h'], 0.7 + 68 / d['page_h'])

    def test_square_card_uses_the_full_width(self):
        sc, ctx, _ = self.square_page(focus=[0.06, 0.315, 0.86, 0.06])
        self.assertEqual(scenes.doc_frame(sc, ctx)['card_w'], 920)

    def test_explicit_crop_is_shown_whole(self):
        sc, ctx, _ = self.square_page(focus=[0.06, 0.315, 0.86, 0.06], crop=[0, 0, 1, 1])
        ctx['crop'] = [0, 0, 1, 1]
        d = scenes.doc_frame(sc, ctx)
        self.assertEqual((d['y'], d['h']), (0, 1))                            # the window scales down to fit instead


class Headlines(unittest.TestCase):
    def lines(self, text, width=920, px=136, max_lines=3):
        html, size = scenes.fit_headline(text, width, px, max_lines=max_lines)
        return [re.sub(r'<[^>]+>', '', line) for line in html.split('<span class="line">') if line], size

    def test_breaks_at_sentence_ends(self):
        lines, size = self.lines("30 days of AI traffic. [Here's what it showed.]")
        for line in lines:
            self.assertNotRegex(line.replace('&#x27;', "'"), r'[.!?] \S', f'a sentence ends mid-line: {lines}')
        self.assertEqual(size, 136)

    def test_balanced_split_kept_when_it_already_respects_sentences(self):
        lines, _ = self.lines('AI moves fast. [Security has to see first.]', 912, 148)
        self.assertEqual(lines, ['AI moves fast.', 'Security has', 'to see first.'])

    def test_square_hook_uses_the_square_hero_scale(self):
        html = scenes.a_hook({'headline': '[Three moves] for safe AI adoption.'},
                             {'id': 'hook', 'index': 0, 'fmt': 'square', 'w': 1080, 'h': 1080, 'pad_x': 80, 'bg': 'pine'})
        self.assertIn('font-size: 136px', html)


class SquareLayout(unittest.TestCase):
    ctx = {'id': 'x', 'index': 1, 'fmt': 'square', 'w': 1080, 'h': 1080, 'pad_x': 80, 'bg': 'pine'}

    def test_compare_sides_share_row_tracks(self):
        html = scenes.a_compare({'before': {'rows': ['a', 'b']}, 'after': {'rows': ['c', 'd', 'e']}}, self.ctx)
        self.assertIn('style="--rows: 4"', html)                                   # title + 3 rows: subgrid tracks in cato.css

    def test_three_cards_use_a_grid(self):
        cards = [{'value': '77', 'label': 'a'}, {'value': '30,254', 'label': 'b'}, {'value': '71%', 'label': 'c'}]
        self.assertIn('class="cards is-grid"', scenes.a_cards({'cards': cards}, self.ctx))
        self.assertIn('class="cards is-column"', scenes.a_cards({'cards': cards[:2]}, self.ctx))

    def test_numbers_size_through_a_variable(self):
        # cato.css scales --fs with the ring (and cards) per format, e.g. a smaller ring above square captions.
        self.assertIn('style="--fs: 260px"', scenes.a_stat({'value': '71%', 'label': 'x'}, self.ctx))
        css = (SKILL / 'assets/kit/cato.css').read_text()
        self.assertIn('calc(var(--fs, var(--t-mega)) * var(--ring-k, 1))', css)
        self.assertIn('.has-captions.fmt-square .stat-ring', css)


class BrandGate(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        sb = {'format': 'portrait', 'scenes': [{'id': 'hook', 'archetype': 'hook', 'duration': 4, 'headline': 'Clear [idea]'},
                                               {'id': 'end', 'archetype': 'end', 'duration': 3, 'line': 'Read more', 'cta': 'Read the report'}]}
        (self.tmp / 'storyboard.json').write_text(json.dumps(sb))
        kit.main(['build', str(self.tmp)])
        self.comp = self.tmp / 'composition'

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def rules(self, findings, severity='error'):
        return {f['rule'] for f in findings if f['severity'] == severity}

    def test_clean_build_passes(self):
        self.assertEqual(self.rules(brand_check.check(self.comp, None)), set())

    def test_violations_are_caught(self):
        f = self.comp / 'compositions/hook.html'
        html = f.read_text()
        html = html.replace('Clear <span class="hl">idea</span>', 'A REVOLUTIONARY, devastating colour shift! Email bob@corp.example TODO')
        html = html.replace('/* Scene-specific layout', '.x { color: #ff00aa; font-family: Comic Sans MS; transform: rotateY(30deg); }\n      /* Scene-specific layout')
        f.write_text(html)
        findings = brand_check.check(self.comp, None)
        errors, warns = self.rules(findings), self.rules(findings, 'warn')
        self.assertTrue({'B-COLOR', 'B-FONT', 'B-ALARM', 'B-PRIVACY', 'B-PLACEHOLDER'} <= errors, errors)
        self.assertTrue({'B-HYPE', 'B-SPELL-US', 'B-CAPS', 'B-MOTION'} <= warns, warns)

    def test_bouncy_ease_rule_reads_code_not_copy(self):
        # Narration in captions.html once tripped B-MOTION: "through" ends in "rough".
        f = self.comp / 'compositions/hook.html'
        html = f.read_text()
        f.write_text(html.replace('CatoMotion.scene(tl, root);', 'CatoMotion.scene(tl, root);\n        const said = "reach AI apps through one policy, a thorough review";'))
        self.assertNotIn('B-MOTION', self.rules(brand_check.check(self.comp, None), 'warn'))
        f.write_text(html.replace('CatoMotion.scene(tl, root);', 'CatoMotion.scene(tl, root);\n        tl.to(root, { y: 0, ease: "elastic.out(1, 0.3)" }, 1);'))
        self.assertIn('B-MOTION', self.rules(brand_check.check(self.comp, None), 'warn'))

    def test_modified_logo_is_rejected(self):
        logo = self.comp / 'assets/cato/logos/cato_logo_rgb_white_art_101.svg'
        logo.write_text(logo.read_text().replace('#fff', '#00FFCC'))
        self.assertIn('B-LOGO', self.rules(brand_check.check(self.comp, None)))

    def test_unknown_product_name_warns(self):
        found = []
        brand_check.copy_rules('x', 'Cato AI Firewall stops prompt injection. Cato SASE Platform and Cato XDR. Cato moves faster.',
                               lambda *a: found.append(a))
        rules = [(f[0], f[3]) for f in found]
        self.assertTrue(any(r == 'B-NAME-UNKNOWN' and 'Cato AI Firewall' in m for r, m in rules), rules)
        self.assertFalse(any(r == 'B-NAME-UNKNOWN' and ('XDR' in m or 'Platform' in m) for r, m in rules), rules)

    def test_product_names_in_post(self):
        post = self.tmp / 'post.md'
        post.write_text('# LinkedIn post\n\nThe Cato SASE Cloud Platform now does more.\n\n#SASE')
        findings = brand_check.check(self.comp, post)
        self.assertIn('B-NAME', self.rules(findings))


class Delivery(unittest.TestCase):
    @unittest.skipUnless(shutil.which('ffmpeg'), 'ffmpeg')
    def test_two_pass_loudness_hits_the_target(self):
        import deliver
        tmp = Path(tempfile.mkdtemp())
        try:
            src, out = tmp / 'bed.wav', tmp / 'out.wav'
            # 25 s of quiet music-like audio: a chord with a slow swell, well below the -20 LUFS target
            subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i',
                            'aevalsrc=0.05*(sin(2*PI*220*t)+sin(2*PI*277*t)+sin(2*PI*330*t))*(0.6+0.4*sin(2*PI*0.2*t)):s=48000:d=25',
                            str(src)], check=True)
            subprocess.run(['ffmpeg', '-v', 'error', '-i', str(src), '-af', deliver.loudnorm(src, -20), '-ar', '48000', str(out)], check=True)
            r = subprocess.run(['ffmpeg', '-hide_banner', '-nostats', '-i', str(out), '-af', 'ebur128', '-f', 'null', '-'], capture_output=True, text=True)
            measured = float(re.findall(r'I:\s+(-?[\d.]+) LUFS', r.stderr)[-1])
            self.assertAlmostEqual(measured, -20, delta=0.5)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class Voice(unittest.TestCase):
    @unittest.skipUnless(shutil.which('say') and shutil.which('ffmpeg'), 'macOS say and ffmpeg')
    def test_say_engine_end_to_end(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            sb = {'format': 'square', 'music': 'none', 'voice': {'engine': 'kokoro', 'voice': 'af_heart'}, 'scenes': [
                {'id': 'hook', 'archetype': 'hook', 'duration': 1.0, 'headline': 'One [idea]', 'say': 'Here is one idea, said slowly enough to need time.'},
                {'id': 'end', 'archetype': 'end', 'duration': 3, 'line': 'Go', 'cta': 'Read it', 'say': ['Read it.', 'Today.']}]}
            (tmp / 'storyboard.json').write_text(json.dumps(sb))
            kit.main(['build', str(tmp)])
            voice.main([str(tmp), '--engine', 'say'])                    # no --voice: af_heart must not carry over
            sb = json.loads((tmp / 'storyboard.json').read_text())
            self.assertEqual((sb['voice']['engine'], sb['voice']['voice']), ('say', 'Samantha'))
            self.assertEqual(sb['scenes'][0]['plannedDuration'], 1.0)      # stretched to fit the line
            self.assertGreater(sb['scenes'][0]['duration'], 2)
            comp = tmp / 'composition'
            index = (comp / 'index.html').read_text()
            self.assertIn('has-captions', index)
            self.assertIn('assets/audio/voiceover.wav', index)
            self.assertIn('CatoMotion.captions', (comp / 'compositions/captions.html').read_text())
            srt = (tmp / 'captions.srt').read_text()
            self.assertIn('Here is one idea', srt)
            self.assertIn('Today.', srt)
            timing = json.loads((tmp / 'voice-timing.json').read_text())
            self.assertEqual([(l['scene'], l['line']) for l in timing['lines']], [('hook', 1), ('end', 1), ('end', 2)])
            import wave
            with wave.open(str(comp / 'assets/audio/voiceover.wav')) as w:
                self.assertAlmostEqual(w.getnframes() / w.getframerate(), json.loads((tmp / 'timing.json').read_text())['duration'], delta=0.05)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    @unittest.skipUnless(shutil.which('say'), 'macOS say')
    def test_unknown_say_voice_is_rejected(self):
        # `say -v <unknown>` exits 0 with the system voice; voice.py must stop instead of narrating in the wrong voice.
        self.assertIn('Samantha', voice.say_voices())
        with self.assertRaises(SystemExit):
            voice.check_voice('say', 'af_heart')
        voice.check_voice('kokoro', 'af_heart')                            # Kokoro names are checked by Kokoro

    def test_spoken_forms(self):
        self.assertEqual(voice.spoken('Cato SASE Platform and SD-WAN for [CISOs]', {}), 'Cato sassy Platform and S D WAN for see-sos')

    def test_caption_cues(self):
        cues = voice.cues_for('One two three four five six seven eight nine ten', 1.0, 4.0)
        self.assertEqual(len(cues), 2)
        self.assertAlmostEqual(cues[-1]['end'], 4.0, places=2)


class Music(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_library_has_ten_original_tracks_with_cues(self):
        originals = [k for k, v in kit.MUSIC.items() if v[2] == 'original']
        self.assertGreaterEqual(len(originals), 10)
        for key, (name, mood, source) in kit.MUSIC.items():
            track = SKILL / 'assets/music' / name
            self.assertTrue(track.is_file() and track.stat().st_size > 100_000, key)
            cue = SKILL / 'assets/music/cues' / name.replace('.mp3', '.music-cues.json')
            doc = json.loads(cue.read_text())
            self.assertGreater(doc['tempo'], 60, key)
            self.assertGreater(len(doc['beats']), 40, key)
            self.assertTrue(doc['strongCues'], key)
            if source == 'original':
                self.assertTrue(doc['sections'], key)

    def build(self, music):
        out = self.tmp / 'v'
        out.mkdir(exist_ok=True)
        (out / 'storyboard.json').write_text(json.dumps({'format': 'square', 'music': music, 'scenes': [
            {'id': 'hook', 'archetype': 'hook', 'duration': 4, 'headline': 'One [idea]'},
            {'id': 'end', 'archetype': 'end', 'duration': 3, 'line': 'Go', 'cta': 'Read it'}]}))
        kit.main(['build', str(out)])
        return out / 'composition'

    def test_new_key_builds_with_cues(self):
        comp = self.build('deep')
        self.assertIn('src="assets/music/robin-deep.mp3"', (comp / 'index.html').read_text())
        self.assertTrue((comp / 'assets/music/cues/robin-deep.music-cues.md').is_file())

    def test_custom_music_file(self):
        import wave as wv, math
        music = self.tmp / 'v/music'
        (music / 'cues').mkdir(parents=True)
        with wv.open(str(music / 'my-bed.wav'), 'wb') as w:
            w.setnchannels(1); w.setsampwidth(2); w.setframerate(8000)
            w.writeframes(b''.join(int(8000 * math.sin(i / 8000 * 2 * math.pi * 220)).to_bytes(2, 'little', signed=True) for i in range(8000 * 3)))
        (music / 'cues/my-bed.music-cues.json').write_text('{"tempo": 100, "beats": [], "strongCues": []}')
        comp = self.build('music/my-bed.wav')
        self.assertIn('src="assets/music/my-bed.wav"', (comp / 'index.html').read_text())
        self.assertTrue((comp / 'assets/music/cues/my-bed.music-cues.json').is_file())

    def test_bad_music_value_is_rejected(self):
        with self.assertRaises(SystemExit):
            kit.validate({'music': 'disco', 'scenes': [{'id': 'a', 'archetype': 'hook', 'duration': 3}]})

    @unittest.skipUnless(shutil.which('ffmpeg'), 'needs ffmpeg')
    def test_composer_renders_a_style(self):
        try:
            import numpy  # noqa: F401
            import scipy  # noqa: F401
        except ImportError:
            self.skipTest('compose_music.py needs numpy and scipy (uv run --with numpy --with scipy)')
        import compose_music
        compose_music.main(['--out', str(self.tmp / 'music'), '--style', 'mallets', '--seed', '3'])
        mp3 = self.tmp / 'music/robin-mallets-s3.mp3'
        self.assertTrue(mp3.is_file() and mp3.stat().st_size > 200_000)
        doc = json.loads((self.tmp / 'music/cues/robin-mallets-s3.music-cues.json').read_text())
        self.assertEqual(doc['tempo'], 108.0)


if __name__ == '__main__':
    unittest.main(verbosity=2)
