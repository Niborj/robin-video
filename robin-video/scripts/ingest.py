"""Robin-Video ingestion: turn source documents into a digest a social video can be planned from.

Reads PDF, PowerPoint (.pptx), Word (.docx), plain text, Markdown, HTML, CSV and images. Writes:

  <out>/content.json   structured blocks, candidate stats, quotes, charts, tables, images, pages
  <out>/content.md     the same as a readable digest with source locations ([p3], [slide 4])
  <out>/pages/         PDF pages (and PPTX slides when LibreOffice is installed) rendered to PNG
  <out>/images/        embedded images extracted from the documents

Usage (from anywhere; uv provisions the dependencies from this folder's pyproject.toml):

  uv run --project <skill-dir>/scripts python <skill-dir>/scripts/ingest.py deck.pptx report.pdf --out <output-dir>/source

Stats and quotes are heuristic candidates. The agent reads content.md in full and verifies every
claim against the source before it goes on screen.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import html.parser
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

TEXT_EXT = {'.txt', '.md', '.markdown', '.rst'}
HTML_EXT = {'.html', '.htm'}
IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.svg'}
CONVERT_EXT = {'.ppt': 'pptx', '.pps': 'pptx', '.ppsx': 'pptx', '.key': 'pdf', '.doc': 'docx', '.rtf': 'docx', '.odt': 'docx', '.odp': 'pptx', '.pages': 'pdf'}
BULLET = re.compile(r'^\s*(?:[•·▪●○◦■□►▶✓✔\-–—*+]|\d{1,2}[.)]|[a-z][.)])\s+')
STAT = re.compile(r"""(?ix)
    (?<![\w.])(
        [$€£¥]\s?\d[\d,]*(?:\.\d+)?\s?(?:k|m|bn|b|tn|t|million|billion|trillion)?\b  # money
      | \d[\d,]*(?:\.\d+)?\s?%                                                       # percent
      | \d+(?:\.\d+)?\s?[x×](?![a-z0-9])                                             # multiplier
      | \d[\d,]*(?:\.\d+)?\+?\s?(?:million|billion|trillion|thousand|k\b|m\b|bn\b|tn\b)  # magnitude
      | \d+\s(?:in|out\sof)\s\d+\b                                                   # 1 in 3
      | \d[\d,]{2,}\+                                                                # 3,000+
      | \d+\+\s(?=[a-z])                                                             # 75+ PoPs
      | \d{1,3}(?:,\d{3})+(?=\s[a-zA-Z])                                            # 30,254 actions
    )""")
CONFIDENTIAL = re.compile(r'\b(confidential|proprietary|internal (?:use )?only|internal assessment|not for (?:external |public )?(?:distribution|release)|do not (?:distribute|share|forward)|under nda|embargo(?:ed)?|draft|pre-release|for internal use)\b', re.I)
INTERNAL_AUDIENCE = re.compile(r'\b(every (?:\w+ )?employee|employees? (?:need|should|must)|bootcamp|enablement|internal training|all[- ]hands|battle ?card|sales playbook|for (?:our )?(?:sales|field) teams?|QBR|internal tool)\b', re.I)
EMAIL = re.compile(r'\b[\w.+-]+@[\w-]+\.[\w.]+\b')
PHONE = re.compile(r'(?<!\d)(?:\+\d{1,3}[\s.-]?)?(?:\(\d{2,4}\)[\s.-]?)?\d{3,4}[\s.-]\d{3,4}(?:[\s.-]\d{2,4})?(?!\d)')
URLISH = re.compile(r'https?://\S+|\b(?:api|internal|intranet|staging|dev)\.[\w.-]+\b', re.I)
QUOTE = re.compile(r'[“"]([^”"]{30,320})[”"]\s*(?:[—–-]\s*([A-Z][^.\n]{2,80}))?')
YEAR = re.compile(r'^(19|20)\d{2}$')


def slug(text, n=40):
    return re.sub(r'[^a-z0-9]+', '-', str(text).lower()).strip('-')[:n].strip('-') or 'x'


def clean(text):
    text = html.unescape(str(text)).replace(' ', ' ').replace('​', '').replace('\r', '')
    text = re.sub(r'[ \t]+', ' ', text)
    return re.sub(r'\s*\n\s*', '\n', text).strip()


def oneline(text):
    return re.sub(r'\s+', ' ', clean(text)).strip()


def words(text):
    return len(str(text).split())


def sentences(text):
    return [s.strip() for s in re.split(r'(?<=[.!?])\s+(?=[A-Z0-9“"])', oneline(text)) if s.strip()]


def _ink(png: Path):
    """A 1-bit-style mask of everything that differs from the page's corner color (light and dark pages alike)."""
    from PIL import Image, ImageChops
    im = Image.open(png).convert('RGB')
    bg = Image.new('RGB', im.size, im.getpixel((2, 2)))
    return ImageChops.difference(im, bg).convert('L').point(lambda v: 255 if v > 18 else 0)


def blank_bands(png: Path, min_frac=0.006):
    """Horizontal bands with no ink, as [y0, y1] fractions of the page height (at least min_frac tall). A doc scene
    snaps its window edges into these so a cropped page never slices through a line of text."""
    try:
        diff = _ink(png)
    except Exception:
        return []
    w, h = diff.size
    data = diff.tobytes()
    bands, start = [], None
    for r in range(h + 1):
        blank = r < h and b'\xff' not in data[r * w:(r + 1) * w]
        if blank and start is None:
            start = r
        elif not blank and start is not None:
            if r - start >= max(2, min_frac * h):
                bands.append([round(start / h, 4), round(r / h, 4)])
            start = None
    return bands


def content_box(png: Path, margin=0.025):
    """Fractional [x, y, w, h] of the non-blank area of a rendered page (so a video can frame the content, not the
    empty margins). Works for light and dark pages by comparing against the corner color."""
    try:
        diff = _ink(png)
        w, h = diff.size
        bbox = diff.getbbox()
        if not bbox:
            return [0, 0, 1, 1]
        x0, y0, x1, y1 = bbox
        x0, y0 = max(0, x0 / w - margin), max(0, y0 / h - margin)
        x1, y1 = min(1, x1 / w + margin), min(1, y1 / h + margin)
        return [round(x0, 4), round(y0, 4), round(x1 - x0, 4), round(y1 - y0, 4)]
    except Exception:
        return [0, 0, 1, 1]


class Digest:
    """Collects blocks and assets from every source in reading order."""

    def __init__(self, out: Path, max_pages: int, page_width: int, min_image: int):
        self.out, self.max_pages, self.page_width, self.min_image = out, max_pages, page_width, min_image
        self.sources, self.blocks, self.images, self.pages, self.charts, self.tables, self.notes = [], [], [], [], [], [], []
        self._hashes = set()
        self.boiler = []   # repeated headers/footers dropped from the text (still scanned for clearance markers)
        (out / 'pages').mkdir(parents=True, exist_ok=True)
        (out / 'images').mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------ recording
    def add(self, src, loc, kind, text, **extra):
        text = clean(text)
        if not text:
            return
        if kind != 'table' and len(text) < 2:
            return
        self.blocks.append({'id': f'b{len(self.blocks) + 1}', 'source': src, 'loc': loc, 'type': kind, 'text': text, **extra})

    def add_image(self, src, loc, data: bytes, ext: str, width=None, height=None, label=''):
        digest = hashlib.sha1(data).hexdigest()
        if digest in self._hashes:
            return None
        if width and height and min(width, height) < self.min_image:
            return None
        self._hashes.add(digest)
        ext = ext.lower().lstrip('.') or 'png'
        if ext in ('jpeg',):
            ext = 'jpg'
        if ext not in ('png', 'jpg', 'gif', 'webp', 'svg'):
            ext = self._to_png(data, ext)
            if not ext:
                return None
            data, ext = ext
        name = f'{slug(loc, 20)}-img{len(self.images) + 1}.{ext}'
        (self.out / 'images' / name).write_bytes(data)
        decorative = bool(width and height and max(width, height) <= 640 and 0.75 <= width / height <= 1.33)
        rec = {'path': f'images/{name}', 'source': src, 'loc': loc, 'width': width, 'height': height, 'label': label, 'iconLike': decorative}
        self.images.append(rec)
        return rec

    def _to_png(self, data, ext):
        try:
            from PIL import Image
            import io
            im = Image.open(io.BytesIO(data))
            buf = io.BytesIO()
            im.convert('RGBA' if im.mode in ('RGBA', 'LA', 'P') else 'RGB').save(buf, 'PNG')
            return buf.getvalue(), 'png'
        except Exception:
            return None

    # ------------------------------------------------------------ PDF
    def pdf(self, path: Path, src: int, label_prefix='p', render=True):
        import pymupdf
        doc = pymupdf.open(path)
        meta_title = oneline((doc.metadata or {}).get('title') or '')
        if re.search(r'untitled|^microsoft (word|powerpoint)|\.(docx?|pptx?|pdf)$', meta_title, re.I):
            meta_title = ''
        n_pages = doc.page_count
        # Pass 1: font statistics and repeated header/footer lines.
        sizes, repeats, page_spans = Counter(), Counter(), []
        for pno in range(n_pages):
            page = doc[pno]
            blocks = page.get_text('dict', sort=True)['blocks']
            page_spans.append(blocks)
            seen = set()
            for b in blocks:
                if b.get('type') != 0:
                    continue
                text = ' '.join(s['text'] for l in b['lines'] for s in l['spans']).strip()
                for l in b['lines']:
                    for s in l['spans']:
                        if s['text'].strip():
                            sizes[round(s['size'] * 2) / 2] += len(s['text'])
                key = re.sub(r'\d+', '#', text.lower())[:80]
                if key and key not in seen:
                    seen.add(key)
                    repeats[key] += 1
        body = sizes.most_common(1)[0][0] if sizes else 11
        boiler = {k for k, c in repeats.items() if n_pages >= 4 and c >= max(3, n_pages * 0.4)}
        self.boiler += sorted(boiler)
        first_title = ''
        # Pass 2: blocks in reading order.
        for pno, blocks in enumerate(page_spans):
            loc = f'{label_prefix}{pno + 1}'
            page_chars = 0
            for b in blocks:
                if b.get('type') != 0:
                    continue
                lines = []
                max_size, bold_chars, chars = 0, 0, 0
                for l in b['lines']:
                    t = ''.join(s['text'] for s in l['spans'])
                    if t.strip():
                        lines.append(t)
                    for s in l['spans']:
                        n = len(s['text'].strip())
                        chars += n
                        max_size = max(max_size, s['size'])
                        if s['flags'] & 16 or re.search(r'bold|black|heavy|semibold', s.get('font', ''), re.I):
                            bold_chars += n
                text = clean('\n'.join(lines))
                if not text:
                    continue
                key = re.sub(r'\d+', '#', oneline(text).lower())[:80]
                if key in boiler or re.fullmatch(r'[\d\s/|.-]+', text):
                    continue
                page_chars += len(text)
                w = words(text)
                bold = chars and bold_chars / chars > 0.6
                if max_size >= body * 1.35 and w <= 22:
                    kind, level = 'heading', 1
                elif (max_size >= body * 1.12 or bold) and w <= 16 and not text.rstrip().endswith(('.', ',', ';')):
                    kind, level = 'heading', 2
                else:
                    kind, level = 'paragraph', None
                if kind == 'heading':
                    self.add(src, loc, kind, oneline(text), level=level, size=round(max_size, 1))
                    if pno == 0 and level == 1 and not first_title:
                        first_title = oneline(text)
                    continue
                # Split bullet lists into separate blocks; otherwise one paragraph per block.
                parts = re.split(r'\n(?=\s*(?:[•·▪●○◦■□►▶✓✔\-–—*]|\d{1,2}[.)])\s)', text)
                for part in parts:
                    if BULLET.match(part):
                        self.add(src, loc, 'bullet', oneline(BULLET.sub('', part, count=1)))
                    else:
                        self.add(src, loc, 'paragraph', oneline(part))
            # Embedded images.
            img_area = 0
            page = doc[pno]
            for info in page.get_images(full=True):
                xref = info[0]
                try:
                    pix = pymupdf.Pixmap(doc, xref)
                    if pix.width < self.min_image or pix.height < self.min_image:
                        continue
                    if pix.n - pix.alpha >= 4 or pix.colorspace is None or pix.colorspace.n not in (1, 3):
                        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                    rects = page.get_image_rects(xref)
                    img_area += sum(r.width * r.height for r in rects)
                    self.add_image(src, loc, pix.tobytes('png'), 'png', pix.width, pix.height)
                except Exception:
                    continue
            if render and pno < self.max_pages:
                zoom = self.page_width / page.rect.width
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
                name = f'pages/{label_prefix.strip() or "p"}-{pno + 1:02d}.png'.replace(' ', '')
                pix.save(self.out / name)
                box = content_box(self.out / name)
                coverage = img_area / max(1, page.rect.width * page.rect.height)
                drawings = len(page.get_drawings()) if hasattr(page, 'get_drawings') else 0
                self.pages.append({'loc': loc, 'image': name, 'source': src, 'width': pix.width, 'height': pix.height,
                                   'textChars': page_chars, 'imageCoverage': round(coverage, 2), 'vectorShapes': drawings,
                                   'visual': coverage > 0.35 or drawings > 60 or page_chars < 250, 'contentBox': box,
                                   'gaps': blank_bands(self.out / name)})
        title = meta_title or first_title
        return {'kind': 'pdf', 'pages': n_pages, 'title': title, 'pagesRendered': min(n_pages, self.max_pages) if render else 0}

    # ------------------------------------------------------------ PowerPoint
    def pptx(self, path: Path, src: int):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
        prs = Presentation(str(path))
        title = oneline(prs.core_properties.title or '')

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(sh.shapes)
                else:
                    yield sh

        for i, slide in enumerate(prs.slides, 1):
            loc = f'slide {i}'
            shapes = list(walk(slide.shapes))
            # Top-to-bottom, left-to-right reading order.
            shapes.sort(key=lambda s: ((s.top or 0) // 20000, s.left or 0))
            slide_title = ''
            for sh in shapes:
                ph = sh.placeholder_format.type if sh.is_placeholder else None
                if sh.has_text_frame:
                    is_title = ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
                    is_sub = ph == PP_PLACEHOLDER.SUBTITLE
                    is_body = ph in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
                    if is_title:
                        t = oneline(sh.text_frame.text)
                        if t:
                            slide_title = slide_title or t
                            self.add(src, loc, 'heading', t, level=1)
                        continue
                    for p in sh.text_frame.paragraphs:
                        t = oneline(''.join(r.text for r in p.runs) or p.text if hasattr(p, 'text') else '')
                        if not t:
                            continue
                        if is_sub:
                            self.add(src, loc, 'heading', t, level=2)
                        elif is_body or p.level > 0:
                            self.add(src, loc, 'bullet', t, level=p.level)
                        else:
                            big = any(r.font.size and r.font.size.pt >= 28 for r in p.runs)
                            self.add(src, loc, 'heading' if big and words(t) <= 12 else 'paragraph', t, **({'level': 2} if big and words(t) <= 12 else {}))
                if getattr(sh, 'has_table', False) and sh.has_table:
                    rows = [[oneline(c.text) for c in r.cells] for r in sh.table.rows]
                    self.tables.append({'source': src, 'loc': loc, 'rows': rows})
                    self.add(src, loc, 'table', '\n'.join(' | '.join(r) for r in rows))
                if getattr(sh, 'has_chart', False) and sh.has_chart:
                    ch = sh.chart
                    try:
                        plot = ch.plots[0]
                        rec = {'source': src, 'loc': loc, 'type': str(ch.chart_type).split('.')[-1].split(' ')[0],
                               'title': oneline(ch.chart_title.text_frame.text) if ch.has_title and ch.chart_title.has_text_frame else '',
                               'categories': [str(c) for c in plot.categories],
                               'series': [{'name': s.name, 'values': [v for v in s.values]} for s in plot.series]}
                        self.charts.append(rec)
                    except Exception as e:  # charts with unusual XML still leave the rest of the deck readable
                        self.charts.append({'source': src, 'loc': loc, 'error': str(e)})
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or (sh.is_placeholder and hasattr(sh, 'image')):
                    try:
                        img = sh.image
                        w, h = img.size if hasattr(img, 'size') else (None, None)
                        self.add_image(src, loc, img.blob, img.ext, w, h, label=getattr(sh, 'name', ''))
                    except Exception:
                        pass
            if slide.has_notes_slide:
                notes = clean(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else '')
                if notes:
                    self.add(src, loc, 'notes', notes)
            if i == 1 and not title:
                title = slide_title
        rendered = self._render_via_office(path, src, 'slide ')
        return {'kind': 'pptx', 'slides': len(prs.slides), 'title': title, 'slideWidth': prs.slide_width, 'slideHeight': prs.slide_height,
                'slidesRendered': rendered}

    def _render_via_office(self, path: Path, src: int, prefix: str):
        office = shutil.which('soffice') or shutil.which('libreoffice') or ('/Applications/LibreOffice.app/Contents/MacOS/soffice' if Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists() else None)
        if not office:
            return 0
        with tempfile.TemporaryDirectory() as tmp:
            try:
                subprocess.run([office, '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(path)], check=True, capture_output=True, timeout=300)
            except Exception:
                return 0
            pdfs = list(Path(tmp).glob('*.pdf'))
            if not pdfs:
                return 0
            import pymupdf
            doc = pymupdf.open(pdfs[0])
            count = 0
            for pno in range(min(doc.page_count, self.max_pages)):
                page = doc[pno]
                zoom = self.page_width / page.rect.width
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
                name = f'pages/slide-{pno + 1:02d}.png'
                pix.save(self.out / name)
                self.pages.append({'loc': f'{prefix}{pno + 1}', 'image': name, 'source': src, 'width': pix.width, 'height': pix.height, 'visual': True,
                                   'contentBox': content_box(self.out / name), 'gaps': blank_bands(self.out / name)})
                count += 1
            return count

    # ------------------------------------------------------------ Word
    def docx(self, path: Path, src: int):
        import docx
        from docx.text.paragraph import Paragraph
        d = docx.Document(str(path))
        title = oneline(d.core_properties.title or '')
        state = {'n': 0, 'title': title, 'seen': set()}
        local = lambda el: el.tag.split('}')[-1]

        def para(el):
            p = Paragraph(el, d)
            t = oneline(p.text)
            if not t or (t in state['seen'] and words(t) > 3):
                return
            state['seen'].add(t)
            state['n'] += 1
            loc = f'¶{state["n"]}'
            style = (p.style.name if p.style is not None else '') or ''
            runs = [r for r in p.runs if r.text.strip()]
            bold = bool(runs) and all(r.bold or (r.style is not None and r.style.font.bold) for r in runs)
            pt = max([r.font.size.pt for r in runs if r.font.size] or [0])
            m = re.match(r'Heading (\d)', style)
            if style == 'Title' or (pt >= 20 and words(t) <= 16 and state['n'] <= 4):
                state['title'] = state['title'] or t
                self.add(src, loc, 'heading', t, level=1)
            elif m:
                self.add(src, loc, 'heading', t, level=int(m.group(1)))
            elif 'List' in style or el.find('.//{*}numPr') is not None or BULLET.match(t):
                self.add(src, loc, 'bullet', BULLET.sub('', t, count=1))
            elif 'Quote' in style:
                self.add(src, loc, 'quote', t)
            elif (bold or pt >= 14) and words(t) <= 12 and not t.endswith(('.', ',', ';', ':')):
                self.add(src, loc, 'heading', t, level=2)
            else:
                self.add(src, loc, 'paragraph', t)

        def walk(el):
            for child in el.iterchildren():
                tag = local(child)
                if tag == 'p':
                    para(child)
                    for box in child.iterfind('.//{*}txbxContent'):  # text boxes (Fallback copies dedupe on text)
                        walk(box)
                elif tag == 'tbl':
                    rows = child.findall('{*}tr')
                    cells = [[tc for tc in tr.findall('{*}tc')] for tr in rows]
                    cell_words = [words(''.join(t.text or '' for t in tc.iter('{*}t'))) for r in cells for tc in r]
                    if len(rows) <= 2 or max(cell_words or [0]) > 30:
                        for r in cells:  # layout table: read cells as ordinary content
                            for tc in r:
                                walk(tc)
                    else:
                        data = [[oneline(''.join(t.text or '' for t in tc.iter('{*}t'))) for tc in r] for r in cells]
                        loc = f'¶{state["n"] + 1}'
                        self.tables.append({'source': src, 'loc': loc, 'rows': data})
                        self.add(src, loc, 'table', '\n'.join(' | '.join(r) for r in data))
                elif tag in ('sdt', 'sdtContent', 'customXml', 'smartTag'):
                    walk(child)

        walk(d.element.body)
        title, para_no = state['title'], state['n']
        for rel in d.part.rels.values():
            if 'image' in rel.reltype:
                try:
                    part = rel.target_part
                    self.add_image(src, 'doc', part.blob, Path(part.partname).suffix, *self._dims(part.blob))
                except Exception:
                    pass
        if not title:
            first = next((b for b in self.blocks if b['source'] == src and b['type'] == 'heading'), None)
            title = first['text'] if first else ''
        return {'kind': 'docx', 'paragraphs': para_no, 'title': title}

    @staticmethod
    def _dims(data):
        try:
            from PIL import Image
            import io
            return Image.open(io.BytesIO(data)).size
        except Exception:
            return (None, None)

    # ------------------------------------------------------------ text, Markdown, HTML, CSV, images
    def text(self, path: Path, src: int):
        raw = path.read_text(encoding='utf-8', errors='replace')
        title, para, para_line = '', [], 0

        def flush():
            if para:
                self.add(src, f'line {para_line}', 'paragraph', ' '.join(para))
                para.clear()
        for n, line in enumerate(raw.splitlines(), 1):
            s = line.strip()
            m = re.match(r'^(#{1,6})\s+(.*)', s)
            if not s:
                flush(); continue
            if m:
                flush()
                self.add(src, f'line {n}', 'heading', m.group(2).strip('# '), level=len(m.group(1)))
                title = title or m.group(2).strip('# ')
            elif re.match(r'^(=+|-+)$', s) and para:
                t = ' '.join(para); para.clear()
                self.add(src, f'line {n - 1}', 'heading', t, level=1 if s[0] == '=' else 2)
                title = title or t
            elif s.startswith('>'):
                flush(); self.add(src, f'line {n}', 'quote', s.lstrip('> '))
            elif BULLET.match(s):
                flush(); self.add(src, f'line {n}', 'bullet', BULLET.sub('', s, count=1))
            elif s.startswith('|') and s.endswith('|'):
                flush()
                if not re.fullmatch(r'[|\s:-]+', s):
                    self.add(src, f'line {n}', 'table', ' | '.join(c.strip() for c in s.strip('|').split('|')))
            else:
                if not para:
                    para_line = n
                para.append(re.sub(r'[*_`]{1,3}([^*_`]+)[*_`]{1,3}', r'\1', s))
        flush()
        if not title:
            first = next((b['text'] for b in self.blocks if b['source'] == src), '')
            title = first if words(first) <= 14 else ''
        return {'kind': 'markdown' if path.suffix.lower() in ('.md', '.markdown') else 'text', 'lines': len(raw.splitlines()), 'title': title}

    def html(self, path: Path, src: int):
        digest = self
        raw = path.read_text(encoding='utf-8', errors='replace')

        class P(html.parser.HTMLParser):
            def __init__(self):
                super().__init__()
                self.stack, self.buf, self.title, self.skip = [], [], '', 0

            def handle_starttag(self, tag, attrs):
                if tag in ('script', 'style', 'nav', 'footer', 'noscript', 'svg'):
                    self.skip += 1
                if tag in ('h1', 'h2', 'h3', 'h4', 'p', 'li', 'blockquote', 'figcaption', 'title', 'td', 'th'):
                    self.stack.append(tag); self.buf = []

            def handle_endtag(self, tag):
                if tag in ('script', 'style', 'nav', 'footer', 'noscript', 'svg'):
                    self.skip = max(0, self.skip - 1)
                if self.stack and self.stack[-1] == tag:
                    self.stack.pop()
                    t = oneline(''.join(self.buf)); self.buf = []
                    if not t:
                        return
                    if tag == 'title':
                        self.title = t
                    elif tag[0] == 'h':
                        digest.add(src, 'html', 'heading', t, level=int(tag[1]))
                    elif tag == 'li':
                        digest.add(src, 'html', 'bullet', t)
                    elif tag == 'blockquote':
                        digest.add(src, 'html', 'quote', t)
                    else:
                        digest.add(src, 'html', 'paragraph', t)

            def handle_data(self, data):
                if self.stack and not self.skip:
                    self.buf.append(data)
        p = P(); p.feed(raw)
        h1 = next((b['text'] for b in self.blocks if b['source'] == src and b['type'] == 'heading'), '')
        return {'kind': 'html', 'title': p.title or h1}

    def csv(self, path: Path, src: int):
        with path.open(newline='', encoding='utf-8', errors='replace') as f:
            rows = [r for r in csv.reader(f)][:60]
        self.tables.append({'source': src, 'loc': 'csv', 'rows': rows})
        self.add(src, 'csv', 'table', '\n'.join(' | '.join(r) for r in rows))
        return {'kind': 'csv', 'rows': len(rows), 'title': path.stem.replace('-', ' ').replace('_', ' ')}

    def image(self, path: Path, src: int):
        w, h = self._dims(path.read_bytes()) if path.suffix.lower() != '.svg' else (None, None)
        old = self.min_image
        self.min_image = 0
        self.add_image(src, path.name, path.read_bytes(), path.suffix, w, h, label=path.stem)
        self.min_image = old
        return {'kind': 'image', 'title': path.stem.replace('-', ' ').replace('_', ' ')}

    # ------------------------------------------------------------ analysis
    def stats(self):
        out, seen = [], set()
        for b in self.blocks:
            if b['type'] in ('table',):
                continue
            for sent in sentences(b['text']) if words(b['text']) > 4 else [b['text']]:
                for m in STAT.finditer(sent):
                    value = m.group(1).strip()
                    if YEAR.match(value.rstrip('+')):
                        continue
                    key = (value.lower(), sent[:80])
                    if key in seen:
                        continue
                    seen.add(key)
                    # Figures inside real sentences are usable claims; bare table cells and 0%/100% rarely are.
                    prose = words(sent) >= 8 and len(re.findall(r'\d', sent)) / max(1, len(sent)) < 0.15
                    trivial = re.fullmatch(r'(0|100)(\.0+)?\s?%', value) is not None
                    out.append({'value': value, 'context': sent if words(sent) <= 60 else ' '.join(sent.split()[:60]) + '…', 'loc': b['loc'], 'block': b['id'],
                                'source': b['source'], 'strength': 'strong' if prose and not trivial else 'weak'})
        out.sort(key=lambda s: s['strength'] != 'strong')
        return out

    def flags(self):
        """Things a human must clear before anything from these sources is published. Confidentiality markers count
        where they act as labels (short blocks, headers, footers, banners), not as words inside body sentences."""
        conf, pii, hosts, audience = {}, [], [], {}
        for b in self.blocks + [{'text': t, 'loc': 'header/footer'} for t in self.boiler]:
            text = b['text']
            label_like = len(text.split()) <= 12 or b['loc'] == 'header/footer'
            for m in CONFIDENTIAL.finditer(text):
                word = m.group(1).lower()
                if word == 'draft':   # a status label ("DRAFT", "Draft v2", "(draft)"), not the verb in "auto-draft reports"
                    if not (m.group(0) == 'DRAFT' or re.search(r'^\s*draft\b|\bdraft\s*(?:version|copy|v\d|\d)|\(draft\)|\bdraft\s*[-–—:]', text, re.I)):
                        continue
                if label_like or m.group(0).isupper() or re.search(r'confidential and proprietary|internal use only', text, re.I):
                    conf.setdefault(word, []).append(b['loc'])
            for m in INTERNAL_AUDIENCE.finditer(text):
                audience.setdefault(m.group(1).lower(), []).append(b['loc'])
            for m in EMAIL.finditer(text):
                pii.append({'kind': 'email', 'value': m.group(0), 'loc': b['loc']})
            for m in URLISH.finditer(text):
                if re.search(r'api\.|internal|intranet|staging|dev\.', m.group(0), re.I):
                    hosts.append({'value': m.group(0), 'loc': b['loc']})
        trim = lambda d: {k: sorted(set(v), key=lambda x: (len(x), x))[:12] for k, v in d.items()}
        return {'confidentialMarkers': trim(conf), 'internalAudience': trim(audience), 'emails': pii[:20], 'internalHosts': hosts[:20]}

    def quotes(self):
        out = []
        for i, b in enumerate(self.blocks):
            if b['type'] == 'quote':
                nxt = self.blocks[i + 1]['text'] if i + 1 < len(self.blocks) else ''
                attr = nxt.lstrip('—–- ') if re.match(r'^[—–-]\s*\w', nxt) and words(nxt) <= 14 else ''
                out.append({'text': b['text'].strip('“”"'), 'attribution': attr, 'loc': b['loc'], 'block': b['id']})
                continue
            for m in QUOTE.finditer(b['text']):
                if words(m.group(1)) >= 6:
                    out.append({'text': m.group(1).strip(), 'attribution': (m.group(2) or '').strip(), 'loc': b['loc'], 'block': b['id']})
        for q in out:
            q['strength'] = 'strong' if q['attribution'] else 'weak'   # unattributed quoted strings are often examples, not quotes
        return out


def convert_with_office(path: Path, target: str, tmp: Path):
    office = shutil.which('soffice') or shutil.which('libreoffice')
    if not office and Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists():
        office = '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    if not office:
        return None
    subprocess.run([office, '--headless', '--convert-to', target, '--outdir', str(tmp), str(path)], check=True, capture_output=True, timeout=300)
    found = list(tmp.glob(f'{path.stem}.{target}'))
    return found[0] if found else None


def render_markdown(d: dict) -> str:
    L = [f'# Source digest: {d["title"] or "(untitled)"}', '']
    for i, s in enumerate(d['sources']):
        extent = ', '.join(f'{v} {k}' for k, v in s.items() if k in ('pages', 'slides', 'paragraphs', 'lines', 'rows'))
        L.append(f'- Source {i}: `{s["name"]}` ({s["kind"]}{", " + extent if extent else ""}){" — title: " + s["title"] if s.get("title") else ""}')
    L += [f'- Words: {d["wordCount"]:,}', f'- Rendered pages: {len(d["pages"])} (in `pages/`)', f'- Extracted images: {len(d["images"])} (in `images/`)']
    for note in d['notes']:
        L.append(f'- Note: {note}')
    f = d.get('flags') or {}
    if f.get('confidentialMarkers') or f.get('internalAudience') or f.get('emails') or f.get('internalHosts'):
        L += ['', '## Clearance flags (resolve before publishing anything)', '']
        for k, locs in f.get('confidentialMarkers', {}).items():
            L.append(f'- Marked “{k}” at {", ".join(locs)}. Confirm with the owner that the content you use is cleared for public social media.')
        for k, locs in f.get('internalAudience', {}).items():
            L.append(f'- Internal-audience signal “{k}” at {", ".join(locs)}: written for employees or the field, so check what may go public.')
        for label, items in (('Email addresses', f.get('emails', [])), ('Internal hosts or API URLs', f.get('internalHosts', []))):
            if items:
                L.append(f'- {label} ({len(items)}) at {", ".join(sorted({i["loc"] for i in items}))}: never show them on screen.')
        L.append('- Customer names, user names and account data in assessments or reports are confidential unless the owner says otherwise: anonymize them.')
    L += ['', '> Stats and quotes below are machine-found candidates. Verify each against the full text before it goes on screen.', '']
    if d['stats']:
        strong = [s for s in d['stats'] if s['strength'] == 'strong']
        weak = [s for s in d['stats'] if s['strength'] != 'strong']
        L += ['## Candidate stats', '']
        for s in strong[:50]:
            L.append(f'- **{s["value"]}** — {s["context"]} [{s["loc"]}]')
        if not strong:
            L.append('- No figures stated in full sentences; see the table-like figures below.')
        if weak:
            locs = sorted({s['loc'] for s in weak}, key=lambda x: (len(x), x))
            L.append(f'- Plus {len(weak)} table-like or trivial figures (in content.json, `strength: weak`) at {", ".join(locs[:12])}{"…" if len(locs) > 12 else ""}. Read the tables and page renders before using them.')
        L.append('')
    if d['quotes']:
        attributed = [q for q in d['quotes'] if q['strength'] == 'strong']
        L += ['## Candidate quotes', '']
        for q in attributed[:20]:
            L.append(f'- “{q["text"]}” — {q["attribution"]} [{q["loc"]}]')
        rest = len(d['quotes']) - len(attributed)
        if rest:
            L.append(f'- {rest} unattributed quoted passage(s) (in content.json, `strength: weak`): often examples or prompts rather than quotes; they also appear in the full text below.')
        L.append('')
    if d['charts']:
        L += ['## Charts (native data from slides; animate these instead of screenshotting them)', '']
        for c in d['charts']:
            if c.get('error'):
                L.append(f'- [{c["loc"]}] chart could not be read: {c["error"]}')
                continue
            L.append(f'- [{c["loc"]}] {c["type"]} chart{": " + c["title"] if c["title"] else ""} — categories {c["categories"]}')
            for s in c['series']:
                L.append(f'  - {s["name"]}: {s["values"]}')
        L.append('')
    if d['tables']:
        L += ['## Tables', '']
        for t in d['tables'][:12]:
            L.append(f'- [{t["loc"]}] {len(t["rows"])} rows; header: {t["rows"][0] if t["rows"] else []}')
        L.append('')
    if d['pages']:
        vis = [p for p in d['pages'] if p.get('visual')]
        L += ['## Page renders', '', 'Use real pages as "show the thing" frames (cover, a chart page, a key diagram). Pages flagged visual are image- or chart-heavy.', '']
        for p in d['pages']:
            cb = p.get('contentBox') or [0, 0, 1, 1]
            trim = f' — content occupies {cb[2] * 100:.0f}%×{cb[3] * 100:.0f}% (doc scenes crop to it)' if cb[2] * cb[3] < 0.85 else ''
            L.append(f'- `{p["image"]}` [{p["loc"]}] {p["width"]}×{p["height"]}{" — visual" if p.get("visual") else ""}{trim}')
        L.append('')
    if d['images']:
        L += ['## Extracted images', '']
        for im in d['images'][:40]:
            L.append(f'- `{im["path"]}` [{im["loc"]}] {im["width"]}×{im["height"]}{" " + im["label"] if im.get("label") else ""}'
                     f'{" — small and square: likely a stock icon (prefer the Cato outline icons)" if im.get("iconLike") else ""}')
        L.append('')
    L += ['## Full text (reading order)', '']
    last_loc = None
    for b in d['blocks']:
        tag = f'[{b["loc"]}]'
        if b['type'] == 'heading':
            L += ['', f'{"#" * min(6, 2 + (b.get("level") or 1))} {b["text"]} {tag}', '']
        elif b['type'] == 'bullet':
            L.append(f'{"  " * int(b.get("level") or 0)}- {b["text"]}{"" if last_loc == b["loc"] else " " + tag}')
        elif b['type'] == 'quote':
            L += [f'> {b["text"]} {tag}', '']
        elif b['type'] == 'notes':
            L += [f'_Speaker notes {tag}:_ {oneline(b["text"])}', '']
        elif b['type'] == 'table':
            L += [f'Table {tag}:', '```', b['text'], '```', '']
        else:
            L += [f'{b["text"]} {tag}', '']
        last_loc = b['loc']
    return re.sub(r'\n{3,}', '\n\n', '\n'.join(L)).strip() + '\n'


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('files', nargs='+', type=Path)
    ap.add_argument('--out', type=Path, required=True, help='output folder, usually <output-dir>/source')
    ap.add_argument('--max-pages', type=int, default=40, help='render at most this many pages per document')
    ap.add_argument('--page-width', type=int, default=1600, help='width in px of rendered pages')
    ap.add_argument('--min-image', type=int, default=160, help='ignore embedded images smaller than this (px)')
    a = ap.parse_args(argv)
    a.out.mkdir(parents=True, exist_ok=True)
    d = Digest(a.out, a.max_pages, a.page_width, a.min_image)
    notes = []
    with tempfile.TemporaryDirectory() as tmp:
        for f in a.files:
            if not f.is_file():
                sys.exit(f'Not a file: {f}')
            ext = f.suffix.lower()
            src = len(d.sources)
            path = f
            if ext in CONVERT_EXT:
                conv = convert_with_office(f, CONVERT_EXT[ext], Path(tmp))
                if not conv:
                    sys.exit(f'{f.name}: {ext} needs LibreOffice to convert. Export it as PDF (or .pptx/.docx) and ingest that instead.')
                path, ext = conv, conv.suffix.lower()
                notes.append(f'{f.name} converted to {conv.suffix} with LibreOffice for reading.')
            if ext == '.pdf':
                info = d.pdf(path, src)
            elif ext == '.pptx':
                info = d.pptx(path, src)
                if not info['slidesRendered']:
                    notes.append(f'{f.name}: slide renders need LibreOffice (not installed). Recreate slides in HTML from the text, charts and images, or export the deck to PDF and ingest it too for page renders.')
            elif ext == '.docx':
                info = d.docx(path, src)
            elif ext in TEXT_EXT:
                info = d.text(path, src)
            elif ext in HTML_EXT:
                info = d.html(path, src)
            elif ext == '.csv':
                info = d.csv(path, src)
            elif ext in IMAGE_EXT:
                info = d.image(path, src)
            else:
                sys.exit(f'{f.name}: unsupported type {ext}. Supported: PDF, PPTX, DOCX, TXT, MD, HTML, CSV, images (or legacy Office/Keynote/Pages with LibreOffice).')
            d.sources.append({'name': f.name, 'path': str(f.resolve()), **info})
    text_blocks = [b for b in d.blocks if b['type'] != 'table']
    titles = [s['title'] for s in d.sources if s.get('title') and s['kind'] != 'image']
    result = {
        'title': titles[0] if titles else (a.files[0].stem.replace('-', ' ').replace('_', ' ')),
        'sources': d.sources,
        'wordCount': sum(words(b['text']) for b in text_blocks),
        'blocks': d.blocks,
        'stats': d.stats(),
        'quotes': d.quotes(),
        'charts': d.charts,
        'tables': d.tables,
        'images': d.images,
        'pages': d.pages,
        'notes': notes,
        'flags': d.flags(),
    }
    (a.out / 'content.json').write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    (a.out / 'content.md').write_text(render_markdown(result), encoding='utf-8')
    result['slug'] = slug(a.files[0].stem)
    (a.out / 'content.json').write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({'out': str(a.out), 'title': result['title'], 'slug': result['slug'], 'words': result['wordCount'], 'blocks': len(d.blocks),
                      'stats': len(result['stats']), 'quotes': len(result['quotes']), 'charts': len(d.charts),
                      'images': len(d.images), 'pages': len(d.pages), 'notes': notes,
                      'clearance': {k: (list(v) if isinstance(v, dict) else len(v)) for k, v in result['flags'].items() if v}}, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
